"""
Excel Helper Module for Telegram Import/Export (tambah_fitur.md Section 11-18, 34, 43)
Handles Excel template generation, customer data export, parsing, and validation.
"""
import io
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.models import CustomerCreate, CustomerRead
from app.database.repository import CustomerRepository
from app.core.logging import logger

EXCEL_HEADERS = [
    "customer_code",
    "username",
    "customer_name",
    "phone",
    "address",
    "package",
    "monitoring_enabled",
    "notes"
]


def generate_template() -> bytes:
    """
    Generates an Excel template file (.xlsx) with sample data and styling.
    (tambah_fitur.md Section 12 & 13)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Pelanggan"

    # Header styling
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(EXCEL_HEADERS)
    for col_num in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Sample rows (tambah_fitur.md Section 12)
    sample_rows = [
        ["C001", "faiz001", "Faizul", "081234567890", "Kambingan", "10M", 1, "Pelanggan baru"],
        ["C002", "ahmad002", "Ahmad", "082198765432", "Lenteng", "20M", 1, "-"]
    ]

    data_font = Font(name="Calibri", size=11)
    for row_data in sample_rows:
        ws.append(row_data)
        row_num = ws.max_row
        for col_num in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = data_font

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_customers(customers: List[CustomerRead]) -> bytes:
    """
    Exports customer list into an Excel file (.xlsx) matching template structure.
    (tambah_fitur.md Section 14 & 43)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Pelanggan"

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(EXCEL_HEADERS)
    for col_num in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    data_font = Font(name="Calibri", size=11)
    for cust in customers:
        row_data = [
            cust.customer_code or "",
            cust.username or "",
            cust.customer_name or "",
            cust.phone or "",
            cust.address or "",
            cust.package or "",
            1 if cust.monitoring_enabled else 0,
            cust.notes or ""
        ]
        ws.append(row_data)
        row_num = ws.max_row
        for col_num in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = data_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def parse_and_validate_excel(file_bytes: bytes) -> Dict[str, Any]:
    """
    Parses uploaded Excel file, validates schema & data constraints without mutating database.
    Returns preview summary and validated CustomerCreate objects.
    (tambah_fitur.md Section 16, 17 & 18)
    """
    try:
        buffer = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(buffer, data_only=True)
    except Exception as e:
        logger.warning(f"Failed to open Excel file: {e}")
        return {
            "valid": False,
            "error_message": "❌ Format file Excel tidak dapat dibaca atau file rusak."
        }

    ws = wb.active

    # Read header row
    header_row = [str(cell.value or "").strip().lower() for cell in ws[1]]
    
    # Map column positions
    col_map = {}
    for idx, col_name in enumerate(header_row):
        if col_name in EXCEL_HEADERS:
            col_map[col_name] = idx

    if "username" not in col_map or "customer_name" not in col_map:
        return {
            "valid": False,
            "error_message": "❌ Format file Excel tidak sesuai. Kolom 'username' dan 'customer_name' wajib ada."
        }

    seen_usernames = set()
    duplicate_usernames = set()
    errors: List[str] = []
    valid_customers: List[CustomerCreate] = []

    total_data_rows = 0
    new_count = 0
    update_count = 0

    for row_idx in range(2, ws.max_row + 1):
        row_cells = ws[row_idx]
        
        def get_val(key: str) -> str:
            if key in col_map and col_map[key] < len(row_cells):
                val = row_cells[col_map[key]].value
                if val is not None:
                    return str(val).strip()
            return ""

        username = get_val("username")
        customer_name = get_val("customer_name")
        customer_code = get_val("customer_code")
        phone = get_val("phone")
        address = get_val("address")
        package = get_val("package")
        notes = get_val("notes")
        monitoring_raw = get_val("monitoring_enabled")

        # Skip completely empty rows
        if not any([username, customer_name, customer_code, phone, address, package, notes, monitoring_raw]):
            continue

        total_data_rows += 1

        # Validation 1: username empty
        if not username:
            errors.append(f"Baris {row_idx} - username kosong")
            continue

        # Validation 2: customer_name empty
        if not customer_name:
            errors.append(f"Baris {row_idx} - customer_name kosong (username: {username})")
            continue

        # Validation 3: duplicate username within file
        if username in seen_usernames:
            errors.append(f"Baris {row_idx} - username duplikat: {username}")
            duplicate_usernames.add(username)
            continue
        
        seen_usernames.add(username)

        # Parse monitoring_enabled
        monitoring_enabled = True
        if monitoring_raw:
            val_lower = monitoring_raw.lower()
            if val_lower in ("0", "false", "off", "no"):
                monitoring_enabled = False
            elif val_lower in ("1", "true", "on", "yes"):
                monitoring_enabled = True

        cust_obj = CustomerCreate(
            customer_code=customer_code or None,
            username=username,
            customer_name=customer_name,
            phone=phone or None,
            address=address or None,
            package=package or None,
            monitoring_enabled=monitoring_enabled,
            notes=notes or None
        )
        valid_customers.append(cust_obj)

        # Check DB to distinguish new vs update
        existing_db = CustomerRepository.get_customer_by_username(username)
        if existing_db:
            update_count += 1
        else:
            new_count += 1

    return {
        "valid": len(errors) == 0,
        "total_rows": total_data_rows,
        "new_count": new_count,
        "update_count": update_count,
        "error_count": len(errors),
        "errors": errors,
        "valid_customers": valid_customers
    }
