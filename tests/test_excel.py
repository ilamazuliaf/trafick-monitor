"""
Unit Tests for Excel Import/Export and Validation (tambah_fitur.md Section 50)
"""
import io
import pytest
import openpyxl
from app.database.database import init_db
from app.database.models import CustomerCreate, CustomerRead
from app.database.repository import CustomerRepository
from app.telegram.excel import generate_template, export_customers, parse_and_validate_excel


@pytest.fixture(autouse=True)
def setup_test_db(tmp_path, monkeypatch):
    test_db = str(tmp_path / "test_excel.db")
    monkeypatch.setattr("app.core.config.settings.database_path", test_db)
    init_db()
    yield test_db


def test_generate_template():
    template_bytes = generate_template()
    assert len(template_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb.active
    headers = [ws.cell(row=1, column=col).value for col in range(1, 9)]
    assert "username" in headers
    assert "customer_name" in headers


def test_export_customers():
    cust1 = CustomerRead(
        id=1, customer_code="C001", username="faiz001", customer_name="Faizul",
        phone="081234567890", address="Kambingan", package="10M",
        monitoring_enabled=True, notes="-", created_at="2026-10-02", updated_at="2026-10-02"
    )
    export_bytes = export_customers([cust1])
    assert len(export_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(export_bytes))
    ws = wb.active
    assert ws.cell(row=2, column=2).value == "faiz001"
    assert ws.cell(row=2, column=3).value == "Faizul"


def test_parse_and_validate_valid_excel():
    template_bytes = generate_template()
    result = parse_and_validate_excel(template_bytes)

    assert result["valid"] is True
    assert result["error_count"] == 0
    assert result["total_rows"] == 2
    assert len(result["valid_customers"]) == 2


def test_parse_and_validate_missing_username():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["customer_code", "username", "customer_name"])
    ws.append(["C001", "", "Faizul"])  # Missing username

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    result = parse_and_validate_excel(excel_bytes)
    assert result["valid"] is False
    assert result["error_count"] == 1
    assert "username kosong" in result["errors"][0]


def test_parse_and_validate_duplicate_username():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["customer_code", "username", "customer_name"])
    ws.append(["C001", "faiz001", "Faizul 1"])
    ws.append(["C002", "faiz001", "Faizul 2"])  # Duplicate username in file

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    result = parse_and_validate_excel(excel_bytes)
    assert result["valid"] is False
    assert result["error_count"] == 1
    assert "username duplikat" in result["errors"][0]


def test_parse_corrupt_file():
    corrupt_bytes = b"This is not an Excel file"
    result = parse_and_validate_excel(corrupt_bytes)
    assert result["valid"] is False
    assert "Format file Excel tidak dapat dibaca" in result["error_message"]
